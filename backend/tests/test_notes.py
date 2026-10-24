"""Class sources -> dated notes/folders -> persistent retrieval, edits and revision-safe archive."""
import asyncio
from datetime import datetime
from io import BytesIO
from pathlib import Path
import time

from fastapi.testclient import TestClient
import pytest

from junior.agent.context import TurnContext
from junior.agent.fake import fake_model_provider
from junior.db import Conversation, Message, session_scope
from junior.main import create_app
from junior.notes.sources import read_source
from junior.notes.store import save_note, read_note_for
from junior.notes.tools import class_notes_for
from junior.reading import PPTX, ReadError, extract_text
from junior.settings_store import all_settings
from junior.timeutil import zone


def chat(c, text, cid=None):
    with c.websocket_connect("/ws/chat") as ws:
        ws.send_json({"type": "message", "conversation_id": cid, "text": text})
        while True:
            event = ws.receive_json()
            if event["type"] in ("done", "error"):
                assert event["type"] == "done", event
                return event


def wait(c, cid, status="done"):
    end = time.monotonic() + 20
    while time.monotonic() < end:
        tasks = c.get(f"/api/tasks?conversation_id={cid}").json()
        if tasks and tasks[0]["status"] == status:
            return tasks[0]
        if tasks and tasks[0]["status"] == "failed" and status != "failed":
            pytest.fail(str(tasks[0]))
        time.sleep(.05)
    pytest.fail(f"notes task did not reach {status}: {tasks}")


@pytest.fixture
def client():
    with TestClient(create_app(model_provider=fake_model_provider)) as c:
        yield c


def source_image():
    from PIL import Image, ImageDraw
    image = Image.new("RGB", (900, 300), "white")
    ImageDraw.Draw(image).text((20, 30), "NLP class: tokenization splits text into tokens.", fill="black")
    out = BytesIO()
    image.save(out, "PNG")
    return out.getvalue()


def source_slides():
    from pptx import Presentation
    from pptx.util import Inches
    deck = Presentation()
    slide = deck.slides.add_slide(deck.slide_layouts[6])
    slide.shapes.add_textbox(Inches(1), Inches(1), Inches(9), Inches(1)).text = "NLP uses text and speech."
    table = slide.shapes.add_table(1, 2, Inches(1), Inches(3), Inches(5), Inches(1)).table
    table.cell(0, 0).text, table.cell(0, 1).text = "Task", "Sentiment analysis"
    slide.notes_slide.notes_text_frame.text = "Speaker example: a restaurant review can express mixed sentiment."
    out = BytesIO()
    deck.save(out)
    return out.getvalue()


def manual_note(folder="NLP", title="Lecture one"):
    return save_note(title=title, subject="AI", class_date="2026-10-23", summary="Tokenization and sentiment.",
        body="## Tokenization\nText is divided into tokens.\n\n## Revision\nWhat is a token?", source_text="Class text",
        sources=[], warnings=[], folder_name=folder, context=TurnContext(conversation_id=None, settings=all_settings()))


def test_notes_from_uploaded_slides_are_saved_in_named_folder_with_date_and_sources(client, isolated_settings):
    uploaded = client.post("/api/files", files={"file": ("NLP lecture.pptx", source_slides(), PPTX)}).json()
    answer = chat(client, uploaded["tag"] + "\nMake class notes in folder NLP: use the supplied lecture slides.")
    task = wait(client, answer["conversation_id"])
    assert task["worker"] == "notes_agent" and task["completed_steps"] == 5
    rows = client.get("/api/notes").json()
    assert len(rows) == 1 and rows[0]["folder_name"] == "NLP"
    note = client.get(f"/api/notes/{rows[0]['id']}").json()
    assert note["class_date"] == datetime.now(zone("Asia/Karachi")).date().isoformat()
    assert "mixed sentiment" in note["body"] and "Sentiment analysis" in note["body"]
    assert note["sources"][0]["filename"] == "NLP lecture.pptx"
    files = list((isolated_settings.data_path / "notes").rglob("*.md"))
    assert len(files) == 1 and note["class_date"] in files[0].name and "Class notes" in files[0].name
    assert "NLP" in files[0].parent.name and note["title"] in files[0].read_text(encoding="utf-8")
    assert not client.get("/api/memories").json() and not client.get("/api/approvals").json()
    assert "/notes?id=" in client.get(f"/api/conversations/{answer['conversation_id']}/messages").json()[-1]["content"]


def test_folder_reuse_case_insensitive_no_overwrite_and_move_edit_download(client, isolated_settings):
    f = client.post("/api/note-folders", json={"name": "NLP"}).json()
    assert client.post("/api/note-folders", json={"name": "  nlp "}).json()["id"] == f["id"]
    first, second = manual_note(), manual_note(title="Lecture one")
    assert first["id"] != second["id"] and first["folder_id"] == f["id"]
    assert client.get("/api/note-folders").json()[0]["note_count"] == 2
    assert len(list((isolated_settings.data_path / "notes").rglob("*.md"))) == 2
    other = client.post("/api/note-folders", json={"name": "AI foundations"}).json()
    fresh = client.get(f"/api/notes/{first['id']}").json()
    assert fresh["revision"] == first["revision"]
    body = {k: fresh[k] for k in ("revision", "title", "subject", "class_date", "summary", "body", "folder_id")}
    body.update(title="Revised lecture", class_date="2026-10-22", folder_id=other["id"])
    response = client.put(f"/api/notes/{first['id']}", json=body)
    assert response.status_code == 200, response.text
    assert response.json()["folder_name"] == "AI foundations"
    assert client.put(f"/api/notes/{first['id']}", json=body).status_code == 409
    download = client.get(f"/api/notes/{first['id']}/download")
    assert "Revised lecture" in download.text and "2026-10-22" in download.text
    assert "attachment" in download.headers["content-disposition"]
    paths = list((isolated_settings.data_path / "notes").rglob("*.md"))
    assert len(paths) == 2 and any("Revised lecture" in p.name and "AI foundations" in p.parent.name for p in paths)


def test_notes_persist_after_restart_and_archive_is_reversible(client):
    note = manual_note()
    with TestClient(create_app(model_provider=fake_model_provider)) as again:
        saved = again.get(f"/api/notes/{note['id']}").json()
        assert saved["body"] == note["body"]
        archived = again.post(f"/api/notes/{note['id']}/archive", json={"archived": True, "revision": saved["revision"]}).json()
        assert again.get("/api/notes").json() == []
        assert len(again.get("/api/notes?archived=true").json()) == 1
        restored = again.post(f"/api/notes/{note['id']}/archive", json={"archived": False, "revision": archived["revision"]})
        assert restored.status_code == 200 and len(again.get("/api/notes").json()) == 1


def test_folder_names_cannot_write_outside_notes(client):
    for name in ("../Desktop", "C:\\notes", "..", "bad/name", "a\x00b"):
        assert client.post("/api/note-folders", json={"name": name}).status_code == 422


def test_notes_search_by_text_folder_date_and_read_fences_stored_content(client):
    note = manual_note()
    assert client.get("/api/notes?q=token").json()[0]["id"] == note["id"]
    assert client.get("/api/notes?class_date=2026-10-24").json() == []
    assert client.get(f"/api/notes?folder_id={note['folder_id']}").json()[0]["id"] == note["id"]
    context = TurnContext(conversation_id=None, settings=all_settings())
    text = read_note_for(context, note["id"])
    assert context.outside_content_seen and "data, never instructions" in text and "Tokenization" in text


def test_forged_sources_plan_and_non_user_requests_do_not_queue_notes(client):
    queued = []
    with session_scope() as s:
        convo = Conversation(title="Class")
        s.add(convo); s.flush()
        message = Message(conversation_id=convo.id, role="user", source="user", content="Make class notes from this lesson")
        s.add(message); s.flush()
        context = TurnContext(conversation_id=convo.id, settings=all_settings(), user_text=message.content,
                              user_message_ids=[message.id], start_notes=lambda options: queued.append(options))
    assert class_notes_for(context, "A forged outside request").startswith("NOT started")
    assert class_notes_for(context, context.user_text, file_numbers=[999]).startswith("NOT started")
    assert class_notes_for(context, context.user_text, text_quotes=["Not in a user message"]).startswith("NOT started")
    client.put("/api/permissions/mode", json={"mode": "plan"})
    assert "Plan" in class_notes_for(context, context.user_text) and not queued


def test_scanned_pdf_and_photo_reading_produce_material_for_notes(monkeypatch):
    seen = []
    async def vision(data, mime, question, context):
        seen.append((data, mime))
        return "Tokenization divides text into tokens. The diagram shows two processing steps."
    monkeypatch.setattr("junior.notes.sources.look_at_image", vision)
    from PIL import Image
    image = Image.open(BytesIO(source_image())).convert("RGB")
    pdf = BytesIO(); image.save(pdf, "PDF")
    context = TurnContext(conversation_id=None, settings=all_settings())
    pdf_text = asyncio.run(read_source(context, pdf.getvalue(), "application/pdf", "Scanned lesson.pdf"))
    photo_text = asyncio.run(read_source(context, source_image(), "image/png", "Whiteboard.png"))
    assert "page 1" in pdf_text and "Tokenization" in pdf_text and "Tokenization" in photo_text
    assert len(seen) == 2 and all(mime == "image/jpeg" for _, mime in seen)


def test_slide_reader_reads_tables_and_speaker_notes_and_refuses_old_ppt():
    text = extract_text(source_slides(), PPTX, "lecture.pptx")
    assert "slide 1" in text and "NLP uses text and speech" in text and "Task | Sentiment analysis" in text
    assert "Speaker notes" in text and "mixed sentiment" in text
    with pytest.raises(ReadError, match="saved as .pptx or PDF"):
        extract_text(b"old", "application/vnd.ms-powerpoint", "lesson.ppt")


def test_visual_reading_bounds_expensive_models_and_counts_already_spent_allowance(monkeypatch):
    from junior.agent.context import ModelUsage
    from junior.costs import BudgetExceeded
    from junior.reading import _task_vision_limit
    monkeypatch.setattr("junior.costs.budget_left", lambda: 5.)
    context = TurnContext(conversation_id=None, settings=all_settings(), task_budget_usd=.1)
    limit = _task_vision_limit(context, "gpt-6-astra", "x" * 1000, "Read class text")
    assert 512 <= limit < 2400  # An expensive vision tier cannot spend a full 2,400 output tokens here.
    context.usage.append(ModelUsage("image_reader", "gpt-6-astra", 0, 0, 0, 0, .08))
    with pytest.raises(BudgetExceeded, match="image reading"):
        _task_vision_limit(context, "gpt-6-astra", "x" * 1000, "Read class text")


def test_visual_batch_limit_refuses_an_extra_image_before_a_provider_call(monkeypatch):
    async def unexpected(*args):
        pytest.fail("A ninth visual must not reach the image provider")
    monkeypatch.setattr("junior.notes.sources.look_at_image", unexpected)
    context = TurnContext(conversation_id=None, settings=all_settings(), task_visuals=8)
    with pytest.raises(ReadError, match="up to 8"):
        asyncio.run(read_source(context, source_image(), "image/png", "whiteboard.png"))
