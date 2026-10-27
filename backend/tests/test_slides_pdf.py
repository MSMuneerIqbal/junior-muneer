"""Presentation approval/API contracts and real offline PDF rendering."""
import asyncio
from io import BytesIO
import json

from fastapi.testclient import TestClient
import httpx
from pypdf import PdfReader
import pytest

from junior.agent.context import TurnContext
from junior.agent.fake import fake_model_provider
from junior.config import get_settings
from junior.db import ApprovalItem, Conversation, Message, session_scope
from junior.documents.pdf import render_pdf
from junior.documents.pdf_worker import html_document
from junior.documents.schema import ResearchDocument, validate_document
from junior.documents.workspace import queue_workspace
from junior.google import auth, slides
from junior.google.fake import FakeGoogle, SCOPES
from junior.main import create_app
from junior.permissions import set_rule
from junior.settings_store import all_settings
from tests.test_documents import chat, wait

SOURCES = [{"id": 1, "title": "Research reference", "url": "https://example.org/research"}]


def document(count=5):
    return ResearchDocument(title="Introduction to NLP", summary="Language tools and their limits.",
        sections=[{"heading": "What is NLP?", "text": "Natural language processing works with language.\nResults require careful evaluation.", "source_ids": [1]}],
        slides=[{"title": "Introduction to NLP" if i == 0 else f"NLP example {i}",
            "layout": "cover" if i == 0 and count > 2 else "points",
            "bullets": ["Systems work with human language.", "Check errors and limitations."],
            "notes": "Explain the language example and its limitations.", "source_ids": [1]}
            for i in range(count - 1 if count > 1 else 1)], headers=[], rows=[])


@pytest.fixture
def google_client(monkeypatch):
    monkeypatch.setenv("GOOGLE_CLIENT_ID", "client-123")
    monkeypatch.setenv("GOOGLE_CLIENT_SECRET", "synthetic-secret")
    get_settings.cache_clear()
    fake = FakeGoogle()
    monkeypatch.setattr(auth, "TRANSPORT", httpx.MockTransport(fake))
    auth._access.clear()
    with TestClient(create_app(model_provider=fake_model_provider)) as client:
        auth._save("you@example.com", "fake-refresh", SCOPES)
        yield client, fake
    get_settings.cache_clear()


def queue_slides(client, count=5):
    words = f"Make {count} Google Slides about NLP"
    with session_scope() as s:
        convo = Conversation(title="Slides test")
        s.add(convo)
        s.flush()
        msg = Message(conversation_id=convo.id, role="user", source="user", content=words)
        s.add(msg)
        s.flush()
        cid, mid = convo.id, msg.id
    ctx = TurnContext(conversation_id=cid, user_text=words, user_message_ids=[mid], settings=all_settings())
    result = queue_workspace(ctx, document(count), SOURCES, "google_slides", count)
    return result, ctx.approval_ids


def test_slides_are_reviewed_even_with_auto_and_send_once(google_client):
    c, fake = google_client
    set_rule("google.slides.create", "auto")
    result, ids = queue_slides(c)
    assert "waiting for your approval" in result and not fake.presentations
    item = c.get(f"/api/approvals/{ids[0]}").json()
    assert item["from_outside"] and item["status"] == "waiting"
    response = c.post(f"/api/approvals/{ids[0]}/approve")
    assert response.status_code == 200 and response.json()["status"] == "posted", response.text
    deck = next(iter(fake.presentations.values()))
    assert len(deck["slides"]) == 5
    inserts = [r["insertText"] for r in deck["requests"] if "insertText" in r]
    assert len([r for r in inserts if r["objectId"].endswith("_notes")]) == 5
    assert any(SOURCES[0]["url"] in r["text"] for r in inserts)
    assert any(r.get("updateTextStyle", {}).get("style", {}).get("link", {}).get("url") == SOURCES[0]["url"] for r in deck["requests"])
    c.post(f"/api/approvals/{ids[0]}/approve")
    assert len(fake.presentations) == 1


def test_invalid_edited_presentation_never_calls_google(google_client):
    c, fake = google_client
    _, ids = queue_slides(c)
    with session_scope() as s:
        item = s.get(ApprovalItem, ids[0])
        payload = json.loads(item.text)
        payload["document"]["slides"][1]["source_ids"] = [999]
        item.text = json.dumps(payload)
    response = c.post(f"/api/approvals/{ids[0]}/approve")
    assert response.status_code == 200 and not fake.presentations
    assert response.json()["status"] != "posted"


@pytest.mark.parametrize("guard", ["deny", "plan", "pause", "disconnected"])
def test_slides_guard_never_creates_remote_files(google_client, guard):
    c, fake = google_client
    if guard == "deny":
        set_rule("google.slides.create", "deny")
    elif guard == "plan":
        c.put("/api/permissions/mode", json={"mode": "plan"})
    elif guard == "pause":
        c.post("/api/pause", json={"paused": True})
    else:
        auth.forget()
    try:
        _, ids = queue_slides(c)
        for number in ids:
            c.post(f"/api/approvals/{number}/approve")
    except Exception as error:
        # Active task limits refuse Pause/Plan before a draft can be queued.
        assert guard in ("pause", "plan") and "Pause or Plan" in str(error)
    assert not fake.presentations


def test_partial_google_creation_pins_id_and_never_recreates(google_client, monkeypatch):
    c, fake = google_client
    _, ids = queue_slides(c)
    original = slides.call
    async def fail_format(method, url, **kwargs):
        if url.endswith(":batchUpdate"):
            raise auth.GoogleError("Synthetic formatting failure")
        return await original(method, url, **kwargs)
    monkeypatch.setattr(slides, "call", fail_format)
    result = c.post(f"/api/approvals/{ids[0]}/approve")
    assert result.status_code == 200 and result.json()["status"] == "check_manually", result.text
    with session_scope() as s:
        item = s.get(ApprovalItem, ids[0])
        assert item.meta["uploaded_drive_id"] in fake.presentations
        from junior.senders.base import Outgoing
        out = Outgoing(item_id=item.id, action_type=item.action_type, text=item.text, meta=item.meta,
                       platform=item.platform, item_type=item.item_type)
    from junior.senders.google_workspace import GoogleWorkspaceSender
    retried = asyncio.run(GoogleWorkspaceSender().send(out))
    assert retried.outcome == "unclear" and len(fake.presentations) == 1
    from dataclasses import replace
    edited = json.loads(out.text)
    edited["document"]["title"] = "A changed card"
    retried = asyncio.run(GoogleWorkspaceSender().send(replace(out, text=json.dumps(edited),
        meta={**out.meta, "slides_content_complete": True})))
    assert retried.outcome == "unclear" and "changed after" in retried.detail and len(fake.presentations) == 1
    invalid = asyncio.run(GoogleWorkspaceSender().send(replace(out, text="invalid edited content")))
    assert invalid.outcome == "unclear" and "already exists" in invalid.detail and len(fake.presentations) == 1


def test_pause_during_google_create_holds_later_writes(google_client, monkeypatch):
    c, fake = google_client
    _, ids = queue_slides(c)
    original = slides.call
    async def pause_after_create(method, url, **kwargs):
        result = await original(method, url, **kwargs)
        if url == slides.SLIDES and method == "POST":
            c.post("/api/pause", json={"paused": True})
        return result
    monkeypatch.setattr(slides, "call", pause_after_create)
    response = c.post(f"/api/approvals/{ids[0]}/approve")
    assert response.json()["status"] == "check_manually"
    assert len(fake.presentations) == 1 and next(iter(fake.presentations.values()))["requests"] == []


@pytest.mark.parametrize("count", [1, 5, 24])
def test_real_slide_pdf_has_exact_pages_and_links(count):
    raw = asyncio.run(render_pdf(document(count), SOURCES, "slides_pdf", count))
    reader = PdfReader(BytesIO(raw))
    assert len(reader.pages) == count
    assert "Introduction to NLP" in reader.pages[0].extract_text()
    assert any(a.get_object().get("/A", {}).get("/URI") == SOURCES[0]["url"] for a in reader.pages[0].get("/Annots", []))
    if count > 1:
        assert "Sources" in reader.pages[-1].extract_text()
        assert any(a.get_object().get("/A", {}).get("/URI") == SOURCES[0]["url"] for a in reader.pages[-1].get("/Annots", []))


def test_real_report_pdf_contains_report_and_sources():
    raw = asyncio.run(render_pdf(document(), SOURCES, "pdf", 5))
    reader = PdfReader(BytesIO(raw))
    content = "\n".join(p.extract_text() for p in reader.pages)
    assert "Natural language processing" in content and SOURCES[0]["url"] in content


def test_print_content_cannot_inject_html_script_or_network():
    doc = document()
    doc.sections[0].text = '<script>fetch("http://127.0.0.1/private")</script><img src="file:///secret">'
    html = html_document(doc, SOURCES, False, 5)
    assert "<script>" not in html and "<img " not in html
    assert "&lt;script&gt;" in html and "default-src 'none'" in html


def test_pdf_cancellation_stops_owned_worker_and_removes_scratch(monkeypatch):
    from junior.documents import pdf
    from junior.computer import owned_process
    import os
    if os.name != "nt":
        pytest.skip("Windows Job Object ownership")
    record = {}
    class Tree:
        def assign(self, pid):
            record["assigned"] = pid
        def close(self):
            record["closed"] = True
    class Input:
        def write(self, data):
            assert record["assigned"] == 123 and data == b"junior-go\n"
        async def drain(self):
            pass
        def close(self):
            pass
    class Process:
        pid, returncode, stdin = 123, None, Input()
        async def wait(self):
            if self.returncode is None:
                record["waiting"].set()
                await asyncio.Event().wait()
            return self.returncode
        def kill(self):
            self.returncode = -1
    async def spawn(*args, **kwargs):
        record["folder"] = args[-1]
        return Process()
    monkeypatch.setattr(owned_process, "OwnedTree", Tree)
    monkeypatch.setattr(pdf.asyncio, "create_subprocess_exec", spawn)
    async def exercise():
        record["waiting"] = asyncio.Event()
        task = asyncio.create_task(render_pdf(document(), SOURCES, "pdf", 5))
        await record["waiting"].wait()
        task.cancel()
        with pytest.raises(asyncio.CancelledError):
            await task
    asyncio.run(exercise())
    from pathlib import Path
    assert record["closed"] and not Path(record["folder"]).exists()


def test_slides_reject_unreadable_reference_count():
    doc = document(15)
    sources = [{"id": i, "title": "Reference", "url": f"https://example.org/{i}"} for i in range(1, 15)]
    for i, slide in enumerate(doc.slides, 1):
        slide.source_ids = [i]
    with pytest.raises(ValueError, match="at most 12"):
        validate_document(doc, sources, "google_slides", 15)


def test_a_cover_alone_cannot_create_an_unsourced_google_deck():
    doc = document(1)
    doc.slides[0].layout = "cover"
    doc.slides[0].source_ids = []
    with pytest.raises(ValueError, match="at least one cited"):
        validate_document(doc, SOURCES, "google_slides", 1)


def test_chat_pdf_completes_in_origin_and_google_slides_queues(google_client):
    c, fake = google_client
    convo = chat(c, "PDF report: explain natural language processing")["conversation_id"]
    task = wait(c, convo, timeout=60)  # Native print worker has its own 45-second deadline.
    raw = c.get(f"/api/files/{task['file_ids'][0]}/download").content
    assert raw.startswith(b"%PDF-")
    other = chat(c, "Make 5 Google Slides about NLP")["conversation_id"]
    task = wait(c, other)
    assert task["output_format"] == "google_slides" and not fake.presentations
    from pptx import Presentation
    raw = c.get(f"/api/files/{task['file_ids'][0]}/download").content
    assert len(Presentation(BytesIO(raw)).slides) == 5
    assert any(a["action_type"] == "google.slides.create" and a["conversation_id"] == other for a in c.get("/api/approvals").json())
