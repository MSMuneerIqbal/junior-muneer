"""Reading what's inside files: ones you send (WhatsApp, 📎) and ones in your Google Drive (also inside folders)."""

import asyncio
import io

import httpx
import pytest
from fastapi.testclient import TestClient

from junior import files
from junior.agent.context import TurnContext
from junior.agent.tools import TOOL_FACTORIES, read_file_for
from junior.config import get_settings
from junior.google import auth
from junior.google.fake import SCOPES, FakeGoogle, tiny_pdf
from junior.google.tools import drive_folder_for, drive_read_for
from junior.main import create_app
from junior.reading import ReadError, extract_text
from junior.registry import load_registry
from junior.settings_store import all_settings

DOCX = "application/vnd.openxmlformats-officedocument.wordprocessingml.document"
XLSX = "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"


def word_file() -> bytes:
    from docx import Document

    doc = Document()
    doc.add_paragraph("Research plan: an AI tutor for Urdu speakers.")
    table = doc.add_table(rows=1, cols=2)
    table.rows[0].cells[0].text, table.rows[0].cells[1].text = "Week 1", "Survey"
    out = io.BytesIO()
    doc.save(out)
    return out.getvalue()


def excel_file() -> bytes:
    from openpyxl import Workbook

    book = Workbook()
    sheet = book.active
    sheet.title = "Income"
    sheet.append(["Month", "USD"])
    sheet.append(["Oct", 1200])
    out = io.BytesIO()
    book.save(out)
    return out.getvalue()


def test_pdf_word_excel_csv_and_text_are_read():
    assert "Invoice 42: total Rs 5,000" in extract_text(tiny_pdf("Invoice 42: total Rs 5,000"), "application/pdf", "inv.pdf")
    word = extract_text(word_file(), DOCX, "plan.docx")
    assert "AI tutor for Urdu speakers" in word and "Week 1 | Survey" in word
    excel = extract_text(excel_file(), XLSX, "income.xlsx")
    assert "--- sheet 'Income' ---" in excel and "2: Oct | 1200" in excel
    assert extract_text(b"Name,Email\nDr A,a@uni.edu\n", "text/csv", "p.csv") == "1: Name | Email\n2: Dr A | a@uni.edu"
    assert extract_text("سلام، notes".encode(), "text/plain", "n.txt") == "سلام، notes"


def test_files_that_cant_be_read_say_why():
    with pytest.raises(ReadError, match="scanned"):
        extract_text(tiny_pdf(""), "application/pdf", "scan.pdf")
    with pytest.raises(ReadError, match="damaged or password-protected"):
        extract_text(b"not a pdf", "application/pdf", "broken.pdf")
    with pytest.raises(ReadError, match=r"\.docx or PDF"):
        extract_text(b"x", "application/msword", "old.doc")
    with pytest.raises(ReadError, match="can't read this kind"):
        extract_text(b"x", "application/zip", "a.zip")


def test_long_files_are_cut():
    text = extract_text(("word " * 10_000).encode(), "text/plain", "long.txt")
    assert len(text) < 20_200 and text.endswith("(cut here: the file has more than 20,000 characters)")


@pytest.fixture
def app(monkeypatch):
    monkeypatch.setenv("GOOGLE_CLIENT_ID", "client-123")
    monkeypatch.setenv("GOOGLE_CLIENT_SECRET", "secret-xyz")
    monkeypatch.setenv("FAKE_AI", "true")  # photos: the pretend vision model
    get_settings.cache_clear()
    monkeypatch.setattr(auth, "TRANSPORT", httpx.MockTransport(FakeGoogle()))
    auth._access.clear()
    with TestClient(create_app()):
        auth._save("you@example.com", "fake-refresh", SCOPES)
        yield lambda: TurnContext(conversation_id=None, settings=all_settings())
    get_settings.cache_clear()


def run(coro):
    return asyncio.run(coro)


def test_a_file_you_sent_is_read_and_fenced(app):
    context = app()
    number, _ = files.save_received(word_file(), "Plan.docx", DOCX, "whatsapp")
    text = run(read_file_for(context, number))
    assert "START OF THEIR TEXT (the file 'Plan.docx' they sent" in text and "AI tutor for Urdu speakers" in text
    assert context.outside_content_seen  # no memory saving from it this turn


def test_a_photo_you_sent_is_looked_at(app):
    context = app()
    number, _ = files.save_received(b"\xff\xd8\xff fake jpeg", "photo 2026-10-18 2140.jpg", "image/jpeg", "whatsapp")
    text = run(read_file_for(context, number, "what's the total?"))
    assert "Invoice 42, total Rs 5,000" in text and "START OF THEIR TEXT" in text


def test_unknown_or_unreadable_files(app):
    context = app()
    assert run(read_file_for(context, 999)).startswith("There's no file 999")
    number, _ = files.save_received(b"PK zip", "stuff.zip", "application/zip", "chat")
    assert run(read_file_for(context, number)).startswith("Couldn't read it: I can't read this kind of file")


def test_reading_a_drive_file_found_anywhere(app):
    context = app()
    text = run(drive_read_for(context, "https://drive.google.com/file/d/file-invoice/view"))
    assert "Invoice 42: total Rs 5,000" in text and "the file 'Invoice September.pdf' from their Google Drive" in text
    doc = run(drive_read_for(context, "https://docs.google.com/document/d/doc-plan/edit"))
    assert "exported as text/plain" in doc and "AI tutor for Urdu speakers" in doc
    assert "is a folder" in run(drive_read_for(context, "https://drive.google.com/drive/folders/x?id=folder-work"))


def test_whats_in_a_folder(app):
    context = app()
    text = run(drive_folder_for(context, "work"))
    assert "Folder 'Work'" in text and "Invoice September.pdf (PDF, changed" in text
    assert "Muneer Iqbal CV.pdf" not in text  # not in that folder
    assert run(drive_folder_for(context, "holidays")).startswith("No folder with 'holidays'")


def test_junior_muneer_has_the_reading_tools():
    assert "files" in load_registry().get("junior_muneer").tools
    assert [t.name for t in TOOL_FACTORIES["files"]()] == ["read_file"]
    names = [t.name for t in TOOL_FACTORIES["google"]()]
    assert "drive_read" in names and "drive_folder" in names
