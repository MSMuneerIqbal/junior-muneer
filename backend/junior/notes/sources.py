"""Class-material reading, including bounded visual extraction. No web research is added to a lecture."""
import asyncio
from io import BytesIO
from threading import RLock

from junior.documents.limits import retry_read
from junior.reading import DOCX, PPTX, ReadError, extract_text, is_image, look_at_image, slide_shapes

MAX_MATERIAL = 80000
MAX_VISUALS = 8
PDF_RENDER_LOCK = RLock()  # PDFium is not thread-safe; multiple notes jobs share this lock.


def _image_bytes(data: bytes) -> bytes:
    from PIL import Image, ImageOps
    with Image.open(BytesIO(data)) as image:
        image = ImageOps.exif_transpose(image).convert("RGB")
        image.thumbnail((1600, 1600))
        output = BytesIO()
        image.save(output, "JPEG", quality=88)
        return output.getvalue()


async def _vision(context, data, label):
    context.check_task_limits()
    if context.task_visuals >= MAX_VISUALS:
        raise ReadError(f"One notes task can read up to {MAX_VISUALS} images/scanned pages. Split these materials into smaller batches.")
    context.task_visuals += 1
    context.task_model_calls += 1
    image = await asyncio.to_thread(_image_bytes, data)
    return await retry_read(context, lambda: look_at_image(image, "image/jpeg",
        f"Transcribe the class material and describe diagrams, formulas and labels in {label}. Mark unclear text; never guess it.", context))


def _pdf_parts(data):
    from pypdf import PdfReader
    reader = PdfReader(BytesIO(data))
    if reader.is_encrypted:
        raise ReadError("This PDF is password-protected. Send an unlocked copy.")
    if len(reader.pages) > 120:
        raise ReadError("Use up to 120 PDF pages per notes task; split this file into lectures or chapters.")
    parts = []
    for index, page in enumerate(reader.pages):
        text = (page.extract_text() or "").strip()
        visual = not text
        if not visual:
            # Small recurring logos do not need a vision call on every page.
            try:
                visual = any(im.image is None or im.image.width * im.image.height > 250000 for im in page.images)
            except Exception:
                visual = True  # let the PDF renderer handle an image format pypdf cannot decode
        parts.append((index, text, visual))
    return parts


def _render_pdf_page(data, index):
    import pypdfium2 as pdfium
    with PDF_RENDER_LOCK, pdfium.PdfDocument(data) as document:
        page = document[index]
        try:
            width, height = page.get_size()
            bitmap = page.render(scale=min(2.0, 1600 / max(width, height)))
            try:
                image = bitmap.to_pil()
                output = BytesIO()
                image.convert("RGB").save(output, "JPEG", quality=88)
                return output.getvalue()
            finally:
                bitmap.close()
        finally:
            page.close()


def _embedded_images(data, mime, name):
    if mime == PPTX or name.lower().endswith(".pptx"):
        from pptx import Presentation
        deck = Presentation(BytesIO(data))
        images = []
        seen = set()
        for n, slide in enumerate(deck.slides, 1):
            for shape in slide_shapes(slide.shapes):
                if hasattr(shape, "image"):
                    image = shape.image
                    # Repeated logos are read once; keep the slide locator on the source.
                    if image.sha1 not in seen:
                        seen.add(image.sha1)
                        images.append((f"slide {n}, image", image.blob))
        return images
    if mime == DOCX or name.lower().endswith(".docx"):
        from docx import Document
        document = Document(BytesIO(data))
        return [(f"document image {n}", part.blob) for n, part in enumerate(document.part.related_parts.values(), 1)
                if part.content_type.startswith("image/")]
    return []


async def read_source(context, data, mime, name):
    context.check_task_limits()
    try:
        if is_image(mime, name):
            return await _vision(context, data, name)
        if mime == "application/pdf" or name.lower().endswith(".pdf"):
            pages = await asyncio.to_thread(_pdf_parts, data)
            if sum(1 for _, _, visual in pages if visual) + context.task_visuals > MAX_VISUALS:
                raise ReadError(f"This batch needs more than {MAX_VISUALS} visual page readings. Split the PDF into smaller sections.")
            chunks = []
            for index, text, visual in pages:
                context.check_task_limits()
                if visual:
                    image = await asyncio.to_thread(_render_pdf_page, data, index)
                    description = await _vision(context, image, f"{name}, page {index + 1}")
                    text += "\nVisual reading:\n" + description
                chunks.append(f"--- page {index + 1} ---\n{text}")
            text = "\n\n".join(chunks)
        else:
            text = await asyncio.to_thread(extract_text, data, mime, name, max_chars=MAX_MATERIAL)
            images = await asyncio.to_thread(_embedded_images, data, mime, name)
            if len(images) + context.task_visuals > MAX_VISUALS:
                raise ReadError(f"This batch contains more than {MAX_VISUALS} images. Split it into smaller lectures.")
            for label, image in images:
                text += f"\n\nVisual reading ({label}):\n" + await _vision(context, image, f"{name}, {label}")
        if "(cut here: the file has more than" in text or len(text) > MAX_MATERIAL:
            raise ReadError("This material is over 80,000 extracted characters. Split it into smaller lectures so no content is silently skipped.")
        return text
    except ReadError:
        raise
    except Exception as e:
        from junior.background import TaskStopped
        from junior.costs import BudgetExceeded
        if isinstance(e, (TaskStopped, BudgetExceeded)):
            raise
        raise ReadError(f"Couldn't read {name} ({type(e).__name__}). Try a supported, unlocked copy.") from e
