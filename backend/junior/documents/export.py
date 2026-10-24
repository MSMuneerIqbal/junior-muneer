"""Deterministic local exports. Models supply content, never executable code or file paths."""
from io import BytesIO
import re

from junior.documents.schema import ResearchDocument

PPTX_MIME = "application/vnd.openxmlformats-officedocument.presentationml.presentation"
DOCX_MIME = "application/vnd.openxmlformats-officedocument.wordprocessingml.document"
XLSX_MIME = "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"


def markdown(doc: ResearchDocument, sources: list[dict]) -> str:
    parts = [f"# {doc.title}", doc.summary]
    for section in doc.sections:
        refs = " ".join(f"[{i}]" for i in section.source_ids)
        parts.append(f"## {section.heading}\n\n{section.text}\n\nSources: {refs}")
    parts.append("## Sources\n\n" + "\n".join(f"[{s['id']}] {s['title']} — {s['url']}" for s in sources))
    return "\n\n".join(parts) + "\n"


def _pptx(doc: ResearchDocument, sources: list[dict], count: int) -> bytes:
    from pptx import Presentation
    from pptx.dml.color import RGBColor
    from pptx.util import Inches, Pt

    prs = Presentation()
    prs.slide_width, prs.slide_height = Inches(13.333), Inches(7.5)
    ink, muted, blue = "162B45", "516277", "196CE6"
    catalog = {s["id"]: s for s in sources}

    def box(slide, text, x, y, w, h, size=24, color=ink, bold=False):
        shape = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
        frame = shape.text_frame
        frame.word_wrap = True
        frame.margin_left = frame.margin_right = Inches(0.02)
        frame.margin_top = frame.margin_bottom = Inches(0.02)
        for index, line in enumerate(text.splitlines()):
            p = frame.paragraphs[0] if index == 0 else frame.add_paragraph()
            p.text = line
            p.font.name = "Aptos"
            p.font.size, p.font.bold = Pt(size), bold
            p.font.color.rgb = RGBColor.from_string(color)
            p.space_after = Pt(14)
            if re.search(r"[\u0600-\u06ff]", line):
                p._p.get_or_add_pPr().set("rtl", "1")
        return shape

    def base(number, dark=False):
        s = prs.slides.add_slide(prs.slide_layouts[6])
        s.background.fill.solid()
        s.background.fill.fore_color.rgb = RGBColor.from_string(ink if dark else "F8FAFD")
        box(s, f"{number:02d} / {count:02d}", 11.7, 7.08, 1.0, .25, 10, "BCCBE0" if dark else muted)
        return s

    for n, content in enumerate(doc.slides, 1):
        cover = content.layout == "cover" and n == 1 and count > 1
        slide = base(n, cover)
        if cover:
            box(slide, content.title, .85, 1.65, 11.4, 1.75, 42, "FFFFFF", True)
            box(slide, "\n".join(content.bullets[:2]), .9, 4.0, 10.5, 1.75, 23, "CDDDF1")
        else:
            box(slide, content.title, .7, .52, 11.9, 1.1, 32, ink, True)
            if content.layout == "comparison" and len(content.bullets) >= 2:
                split = (len(content.bullets) + 1) // 2
                box(slide, "\n\n".join(content.bullets[:split]), .8, 2.0, 5.5, 4.5, 22)
                box(slide, "\n\n".join(content.bullets[split:]), 7.0, 2.0, 5.5, 4.5, 22)
            else:
                prefix = (lambda i: f"{i + 1}. ") if content.layout == "steps" else (lambda i: "• ")
                # Short blocks preserve readable fonts; schema limits content density.
                for i, bullet in enumerate(content.bullets):
                    box(slide, prefix(i) + bullet, .85, 1.85 + i * .89, 11.6, .82, 22)
            refs = "Sources: " + ", ".join(f"[{i}]" for i in content.source_ids)
            box(slide, refs, .8, 6.9, 10.7, .35, 11, muted)
        notes = slide.notes_slide.notes_text_frame
        notes.text = content.notes + "\n\nSources:\n" + "\n".join(
            f"[{i}] {catalog[i]['title']}\n{catalog[i]['url']}" for i in content.source_ids
        )
    if count > 1:
        slide = base(count)
        box(slide, "Sources", .7, .52, 11.9, 1.1, 32, ink, True)
        used = sorted({i for slide in doc.slides for i in slide.source_ids})
        for n, i in enumerate(used[:12]):
            source = catalog[i]
            x, y = .8 + (n // 6) * 6.1, 1.75 + (n % 6) * .8
            title = source["title"]
            if len(title) > 78:
                title = title[:75].rsplit(" ", 1)[0] + "…"
            shape = box(slide, f"[{i}] {title}", x, y, 5.7, .72, 16, blue)
            shape.text_frame.paragraphs[0].runs[0].hyperlink.address = source["url"]
        slide.notes_slide.notes_text_frame.text = "\n".join(
            f"[{i}] {catalog[i]['title']}\n{catalog[i]['url']}" for i in used
        )
    stream = BytesIO()
    prs.save(stream)
    return stream.getvalue()


def export(doc: ResearchDocument, sources: list[dict], output_format: str, count: int = 10) -> tuple[bytes, str, str]:
    if output_format == "pptx":
        return _pptx(doc, sources, count), ".pptx", PPTX_MIME
    if output_format in ("docx", "google_doc"):
        from docx import Document
        document = Document()
        document.add_heading(doc.title, 0)
        document.add_paragraph(doc.summary)
        for part in doc.sections:
            document.add_heading(part.heading, 1)
            document.add_paragraph(part.text)
            document.add_paragraph("Sources: " + ", ".join(f"[{i}]" for i in part.source_ids))
        document.add_heading("Sources", 1)
        for s in sources:
            document.add_paragraph(f"[{s['id']}] {s['title']}\n{s['url']}")
        stream = BytesIO()
        document.save(stream)
        return stream.getvalue(), ".docx", DOCX_MIME
    if output_format in ("xlsx", "google_sheet"):
        from openpyxl import Workbook
        from openpyxl.styles import Font, PatternFill
        book = Workbook()
        sheet = book.active
        sheet.title = "Research"
        for row in [doc.headers, *doc.rows]:
            sheet.append(row)
            for cell in sheet[sheet.max_row]:
                # Spreadsheet text is data, including values beginning with '='.
                cell.data_type = "s"
        sheet.freeze_panes = "A2"
        sheet.auto_filter.ref = sheet.dimensions
        for cell in sheet[1]:
            cell.font = Font(bold=True, color="FFFFFF")
            cell.fill = PatternFill("solid", fgColor="162B45")
        for col in sheet.columns:
            sheet.column_dimensions[col[0].column_letter].width = min(60, max(18, max(len(str(c.value or "")) for c in col) + 2))
        refs = book.create_sheet("Sources")
        refs.append(["ID", "Title", "URL"])
        for s in sources:
            refs.append([s["id"], s["title"], s["url"]])
            for cell in refs[refs.max_row]:
                cell.data_type = "s"
        stream = BytesIO()
        book.save(stream)
        return stream.getvalue(), ".xlsx", XLSX_MIME
    return markdown(doc, sources).encode("utf-8"), ".md", "text/markdown; charset=utf-8"
